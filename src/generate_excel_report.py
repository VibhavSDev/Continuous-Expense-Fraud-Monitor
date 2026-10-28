import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

print("Generating Advanced Excel Audit Checklist...")

# 1. Load the exceptions generated in Phase 2
try:
    df = pd.read_csv("Audit_Exceptions_Report.csv")
except FileNotFoundError:
    print("Error: Run audit_engine.py first to generate the exceptions report.")
    exit()

# 2. Add columns specifically for the manual audit review process
df['Review Status'] = 'Pending'
df['Auditor Comments'] = ''

# 3. Export to Excel using the openpyxl engine
excel_filename = "Audit_Review_Checklist.xlsx"
writer = pd.ExcelWriter(excel_filename, engine='openpyxl')
df.to_excel(writer, index=False, sheet_name='High Risk Transactions')
writer.close()

# 4. Apply Advanced Formatting with openpyxl
wb = load_workbook(excel_filename)
ws = wb['High Risk Transactions']

# Define styling elements
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Blue
header_font = Font(color="FFFFFF", bold=True)
flag_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")   # Light Red
currency_format = '"$"#,##0.00'

# Format the Header Row, Freeze Top Row, and Add AutoFilters
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

# Iterate through columns to adjust widths and format specific data types
for col in ws.columns:
    max_length = 0
    col_letter = col[0].column_letter
    col_header = col[0].value
    
    for cell in col:
        # Format Currency
        if col_header == 'Transaction Amount' and cell.row > 1:
            cell.number_format = currency_format
            
        # Highlight the Flag columns in red to draw the auditor's eye
        if isinstance(col_header, str) and col_header.startswith('Flag_') and cell.row > 1:
            if cell.value == True:
                cell.fill = flag_fill

        # Calculate optimal column width
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
            
    # Set the column width with a little padding
    adjusted_width = (max_length + 2)
    ws.column_dimensions[col_letter].width = adjusted_width

# Set a wider fixed width for the comments column
ws.column_dimensions[get_column_letter(ws.max_column)].width = 40

wb.save(excel_filename)
print(f"Success! Highly formatted auditor checklist saved as '{excel_filename}'.")